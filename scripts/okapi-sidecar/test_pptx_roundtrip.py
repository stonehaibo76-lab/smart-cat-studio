"""
Manual round-trip test for PPTX extract/merge.

Run from scripts/okapi-sidecar:
  python test_pptx_roundtrip.py
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from pptx_handler import extract_pptx, merge_pptx


def build_sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run1 = p.add_run()
    run1.text = "Hello "
    run1.font.bold = True
    run2 = p.add_run()
    run2.text = "World"
    run2.font.italic = True

    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1.5))
    table = table_shape.table
    table.cell(0, 0).text = "Cell A1"
    table.cell(0, 1).text = "Cell B1"
    table.cell(1, 0).text = "Cell A2"
    table.cell(1, 1).text = "Cell B2"

    notes = slide.notes_slide.notes_text_frame
    notes.text = "Speaker note for slide one."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> int:
    original = build_sample_pptx()
    segments = extract_pptx(original)
    if not segments:
        print("FAIL: no segments extracted")
        return 1

    print(f"Extracted {len(segments)} segments:")
    for seg in segments:
        meta = seg.get("inlineRunMeta") or []
        print(f"  {seg['okapiTuId']}: {seg['source']!r} (runs={len(meta)})")

    units = []
    for seg in segments:
        target = seg["source"]
        if "Hello" in target:
            target = target.replace("Hello", "Nihao")
        elif "World" in target:
            target = target.replace("World", "Shijie")
        elif target.startswith("Cell"):
            target = target.replace("Cell", "Danyuan")
        elif "Speaker" in target:
            target = target.replace("Speaker note", "Yanjiang beizhu")
        units.append(
            {
                "okapiTuId": seg["okapiTuId"],
                "source": seg["source"],
                "target": target,
            }
        )

    merged = merge_pptx(original, units)
    prs = Presentation(io.BytesIO(merged))
    slide = prs.slides[0]
    texts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    texts.append(cell.text.strip())
    if slide.has_notes_slide:
        texts.append(slide.notes_slide.notes_text_frame.text.strip())

    joined = " | ".join(t for t in texts if t)
    print(f"Merged texts: {joined!r}")

    if "Nihao" not in joined or "Danyuan" not in joined or "Yanjiang" not in joined:
        print("FAIL: merged content missing expected translations")
        return 1

    out_path = Path(__file__).resolve().parent / "_test_pptx_roundtrip_out.pptx"
    out_path.write_bytes(merged)
    print(f"OK: wrote {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
